from __future__ import annotations

import asyncio
import base64
import binascii
import csv
import hashlib
import io
import json
import os
from datetime import date, datetime


def _get_program_store(pk):
    from services.lib.program_store import ProgramStore

    return ProgramStore(os.getenv("IF_HEALTH_TABLE_NAME", "if-health"), pk, os.getenv("AWS_REGION", "ca-central-1"))


def _get_glossary_store(pk):
    from services.lib.glossary_store import GlossaryStore

    return GlossaryStore(os.getenv("IF_HEALTH_TABLE_NAME", "if-health"), pk, os.getenv("AWS_REGION", "ca-central-1"))


def _get_import_store(pk):
    from .._import_store import ImportStore

    return ImportStore(os.getenv("IF_HEALTH_TABLE_NAME", "if-health"), pk, os.getenv("AWS_REGION", "ca-central-1"))


def _file_hash(content):
    return f"sha256:{hashlib.sha256(content).hexdigest()[:16]}"


def _extract_xlsx(content):
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    rows = []
    sheets = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        if sheet.max_row is None or sheet.max_row < 2:
            continue
        header_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = [str(cell.value).strip() if cell.value else f"col_{index}" for index, cell in enumerate(header_row)]
        sheet_has_data = False
        for cells in sheet.iter_rows(min_row=2):
            row = {"_sheet": sheet_name}
            has_data = False
            for index, cell in enumerate(cells):
                if index >= len(headers):
                    continue
                value = cell.value
                if isinstance(value, (datetime, date)):
                    value = value.isoformat()
                row[headers[index]] = value
                has_data |= value is not None
            if has_data:
                rows.append(row)
                sheet_has_data = True
        if sheet_has_data:
            sheets.append(sheet_name)
    return rows, ", ".join(sheets) if sheets else "empty"


def _extract_csv(content):
    return [row for row in csv.DictReader(io.StringIO(content.decode("utf-8", errors="replace"))) if any(row.values())]


def _preclassify(rows):
    if not rows:
        return None
    keys = rows[0]
    date_keys = [key for key in keys if "date" in key.lower()]
    for row in rows[:20]:
        for key in date_keys:
            value = row.get(key)
            if isinstance(value, (date, datetime)) or isinstance(value, str) and "-" in value and len(value) >= 8:
                return "session_import"
    if any("week" in key.lower() for key in keys):
        return "template"
    if any(any(token in key.lower() for token in ("rpe", "percentage", "%", "target")) for key in keys):
        return "template"
    return None


async def _dispatch(args):
    from .import_parse_ai import generate_import_parse_report

    try:
        content = base64.b64decode(args["base64_content"], validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ValueError("base64_content is invalid") from exc
    filename = args["filename"]
    rows, _ = _extract_xlsx(content) if filename.lower().endswith(".xlsx") else (_extract_csv(content), "")
    classification = _preclassify(rows) or "session_import"
    pk = args["pk"]
    try:
        program = await _get_program_store(pk).get_program()
        context = {
            "current_maxes": program.get("current_maxes", {}),
            "current_program_weeks": len({session.get("week_number") for session in program.get("sessions", []) if session.get("week_number")}),
        }
    except Exception:
        context = {"current_maxes": {}, "current_program_weeks": 0}
    parsed = await generate_import_parse_report(json.dumps(rows[:100], indent=2, default=str), filename, classification, context)
    glossary = _get_glossary_store(pk)
    await glossary.get_glossary()
    names = {
        exercise.get("name")
        for session in parsed.get("sessions", [])
        for exercise in session.get("exercises", [])
        if exercise.get("name")
    }
    resolved = {}
    for name in names:
        glossary_id = await glossary.fuzzy_resolve(name, threshold=0.92)
        if glossary_id:
            resolved[name] = glossary_id
    for session in parsed.get("sessions", []):
        for exercise in session.get("exercises", []):
            if exercise.get("name") in resolved:
                exercise["glossary_id"] = resolved[exercise["name"]]
    import_id = await _get_import_store(pk).stage_import({
        "import_type": "template" if classification == "template" else "session_import",
        "source_filename": filename,
        "source_file_hash": _file_hash(content),
        "ai_parse_result": parsed,
    })
    return {
        "import_id": import_id,
        "classification": classification,
        "warnings": parsed.get("warnings", []),
        "parse_notes": parsed.get("parse_notes", ""),
    }


def import_parse_file(args):
    return asyncio.run(_dispatch(args))
